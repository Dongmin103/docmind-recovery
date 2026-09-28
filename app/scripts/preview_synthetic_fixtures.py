"""Create non-sensitive preview fixtures outside the repository for manual QA."""

import sys
from pathlib import Path

import xlwt
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches


def main(output: Path):
    output.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("임시 원문 미리보기 합성 시험", 0)
    for page in range(3):
        document.add_heading(f"{page + 1}. 한국어 문서", 1)
        for paragraph in range(8):
            document.add_paragraph(f"검증용으로 생성한 문장 {paragraph + 1}. 실제 사용자 자료가 아닙니다.")
        table = document.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "항목"
        table.cell(0, 1).text = "수량"
        table.cell(1, 0).text = "미리보기"
        table.cell(1, 1).text = "3"
        if page < 2:
            document.add_page_break()
    document.save(output / "synthetic.docx")

    slides = Presentation()
    for index in range(3):
        slide = slides.slides.add_slide(slides.slide_layouts[5])
        slide.shapes.title.text = f"한국어 슬라이드 {index + 1}"
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(7), Inches(1))
        box.text_frame.text = "임시 원문 미리보기 · 합성 자료"
        table = slide.shapes.add_table(3, 2, Inches(1), Inches(3), Inches(6), Inches(2)).table
        table.cell(0, 0).text = "항목"
        table.cell(0, 1).text = "값"
        table.cell(1, 0).text = "슬라이드"
        table.cell(1, 1).text = str(index + 1)
    slides.save(output / "synthetic.pptx")

    book = Workbook()
    for index, sheet_name in enumerate(("매출", "합계")):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = sheet_name
        sheet.append(["품목", "수량", "단가"])
        for row in range(1, 21):
            sheet.append([f"합성 항목 {row}", row, row * 100])
        sheet.freeze_panes = "A2"
    book.save(output / "synthetic.xlsx")
    binary_book = xlwt.Workbook()
    for sheet_name in ("매출", "합계"):
        sheet = binary_book.add_sheet(sheet_name)
        for row in range(20):
            sheet.write(row, 0, f"합성 항목 {row + 1}")
            sheet.write(row, 1, row + 1)
    binary_book.save(str(output / "synthetic.xls"))
    print("Created four synthetic Office fixtures; no source documents read.")


if __name__ == "__main__":
    main(Path(sys.argv[1]))
