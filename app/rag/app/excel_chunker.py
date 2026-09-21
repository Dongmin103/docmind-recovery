"""Excel-only chunk boundaries over parser cells; cell values are indivisible."""
from datetime import date, datetime, time
from hashlib import sha256
from html import escape
from io import BytesIO
from pathlib import Path
import re
from common.token_utils import num_tokens_from_string


def excel_embedding_counter(parser_config):
    """Load the explicitly configured local embedding tokenizer without a fallback."""
    budget = parser_config.get("excel_embedding_token_budget")
    if isinstance(budget, bool) or not isinstance(budget, int) or budget < 1:
        raise ValueError("excel_embedding_token_budget must be a positive integer")
    config = parser_config.get("excel_embedding_tokenizer")
    if not isinstance(config, dict):
        raise ValueError("excel_embedding_tokenizer configuration is required")
    path = config.get("path")
    expected_hash = config.get("sha256")
    if not isinstance(path, str) or not path:
        raise ValueError("excel_embedding_tokenizer.path is required")
    if not isinstance(expected_hash, str) or len(expected_hash) != 64:
        raise ValueError("excel_embedding_tokenizer.sha256 is required")
    tokenizer_path = Path(path)
    if not tokenizer_path.is_file():
        raise ValueError("excel_embedding_tokenizer.path is not a readable file")
    actual_hash = sha256(tokenizer_path.read_bytes()).hexdigest()
    if actual_hash != expected_hash:
        raise ValueError("excel_embedding_tokenizer.sha256 does not match the tokenizer file")
    if config.get("add_special_tokens") is not True:
        raise ValueError("excel_embedding_tokenizer.add_special_tokens must be true")
    try:
        from tokenizers import Tokenizer
    except ImportError as exc:
        raise RuntimeError("tokenizers is required for the configured Excel embedding tokenizer") from exc
    tokenizer = Tokenizer.from_file(str(tokenizer_path))
    tokenizer.no_truncation()
    tokenizer.no_padding()
    return budget, lambda text: len(tokenizer.encode(text, add_special_tokens=True).ids)

def _a1(row, column):
    letters = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters
    return f"{letters}{row}"


def _display_value(value, number_format=""):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and (not number_format or number_format.lower() == "general") and value.is_integer():
        return str(int(value))
    if isinstance(value, float) and number_format:
        primary = number_format.split(";")[0]
        if "%" in primary:
            decimals = len(primary.rsplit(".", 1)[1].split("%", 1)[0]) if "." in primary else 0
            return f"{value * 100:.{decimals}f}%"
        if any(token in primary for token in ("0", "#")):
            decimals = len(primary.rsplit(".", 1)[1]) if "." in primary else 0
            rendered = f"{value:,.{decimals}f}" if "," in primary else f"{value:.{decimals}f}"
            return rendered
    return str(value)


def _raw_type(value):
    if value is None:
        return "empty"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, (datetime, date, time)):
        return "date"
    return "text"


def _grid_sheet(name, max_row, max_column, cells, merged_ranges):
    return {
        "name": name,
        "max_row": max_row,
        "max_column": max_column,
        "merged_ranges": sorted(merged_ranges, key=lambda item: (item["min_row"], item["min_column"])),
        "cells": cells,
    }


def _xlsx_canonical_grid(binary):
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(binary), data_only=False, read_only=False)
    cached_workbook = load_workbook(BytesIO(binary), data_only=True, read_only=False)
    sheets = []
    for worksheet in workbook.worksheets:
        cached_worksheet = cached_workbook[worksheet.title]
        cells = []
        for row in range(1, worksheet.max_row + 1):
            for column in range(1, worksheet.max_column + 1):
                cell = worksheet.cell(row=row, column=column)
                cached = cached_worksheet.cell(row=row, column=column)
                raw_type = "formula" if cell.data_type == "f" else _raw_type(cell.value)
                cells.append({
                    "row": row,
                    "column": column,
                    "a1": cell.coordinate,
                    "display": _display_value(cached.value, cell.number_format),
                    "raw_value": _display_value(cell.value),
                    "raw_type": raw_type,
                    "number_format": cell.number_format,
                })
        merges = [
            {
                "min_row": merged.min_row,
                "max_row": merged.max_row,
                "min_column": merged.min_col,
                "max_column": merged.max_col,
            }
            for merged in worksheet.merged_cells.ranges
        ]
        sheets.append(_grid_sheet(worksheet.title, worksheet.max_row, worksheet.max_column, cells, merges))
    return sheets


def _xls_number_format(book, cell):
    if cell.xf_index is None:
        return ""
    key = book.xf_list[cell.xf_index].format_key
    return book.format_map[key].format_str if key in book.format_map else ""


def _xls_display_value(book, cell):
    import xlrd

    if cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
        return ""
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return "True" if cell.value else "False"
    number_format = _xls_number_format(book, cell)
    if cell.ctype == xlrd.XL_CELL_DATE:
        value = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
        lowered = number_format.lower()
        if lowered == "m/d/yy":
            return f"{value.month}/{value.day}/{value.year % 100:02d}"
        if lowered == "mmm-yy":
            return value.strftime("%b-%y")
        if lowered == "h:mm":
            return f"{value.hour}:{value.minute:02d}"
        return value.isoformat()
    return _display_value(cell.value, number_format)


def _xls_raw_type(cell):
    import xlrd

    return {
        xlrd.XL_CELL_EMPTY: "empty",
        xlrd.XL_CELL_BLANK: "empty",
        xlrd.XL_CELL_TEXT: "text",
        xlrd.XL_CELL_NUMBER: "number",
        xlrd.XL_CELL_DATE: "date",
        xlrd.XL_CELL_BOOLEAN: "boolean",
        xlrd.XL_CELL_ERROR: "error",
    }.get(cell.ctype, "unknown")


def _xls_canonical_grid(binary):
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("xlrd with formatting_info support is required for original .xls grids") from exc

    workbook = xlrd.open_workbook(file_contents=binary, formatting_info=True)
    sheets = []
    for worksheet in workbook.sheets():
        cells = []
        for row in range(worksheet.nrows):
            for column in range(worksheet.ncols):
                cell = worksheet.cell(row, column)
                cells.append({
                    "row": row + 1,
                    "column": column + 1,
                    "a1": _a1(row + 1, column + 1),
                    "display": _xls_display_value(workbook, cell),
                    "raw_value": str(cell.value),
                    "raw_type": _xls_raw_type(cell),
                    "number_format": _xls_number_format(workbook, cell),
                })
        merges = [
            {"min_row": first_row + 1, "max_row": last_row, "min_column": first_column + 1, "max_column": last_column}
            for first_row, last_row, first_column, last_column in worksheet.merged_cells
        ]
        sheets.append(_grid_sheet(worksheet.name, worksheet.nrows, worksheet.ncols, cells, merges))
    return sheets


def load_excel_canonical_grid(binary, filename):
    """Read an original workbook into an immutable, coordinate-complete grid."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        sheets = _xlsx_canonical_grid(binary)
    elif name.endswith(".xls"):
        sheets = _xls_canonical_grid(binary)
    else:
        raise ValueError(f"unsupported Excel grid format: {filename}")
    return {"schema_version": "excel-grid-v1", "sheets": sheets}


def render_excel_display_html(grid, source_cells, *, compact=True):
    """Render a compact, slot-complete projection of explicitly selected cells.

    The displayed rows are the selected source rows.  Every rendered row keeps
    its source row number and every selected horizontal extent has explicit
    empty slots, so merged spans remain verifiable without disclosing values
    from unselected canonical cells.
    """
    sheets = {sheet["name"]: sheet for sheet in grid.get("sheets", [])}
    selected_by_sheet = {}
    for cell in source_cells:
        selected_by_sheet.setdefault(cell["sheet"], []).append(cell)
    if not selected_by_sheet:
        raise ValueError("Excel display HTML requires at least one selected source cell")
    if len(selected_by_sheet) > 1:
        return "".join(
            render_excel_display_html({"sheets": [sheets[sheet_name]]}, selected_by_sheet[sheet_name], compact=compact)
            for sheet_name in sheets if sheet_name in selected_by_sheet
        )
    sheet_name, selected = next(iter(selected_by_sheet.items()))
    sheet = sheets.get(sheet_name)
    if sheet is None:
        raise ValueError(f"Excel display grid is missing sheet {sheet_name!r}")

    cell_by_position = {(cell["row"], cell["column"]): cell for cell in sheet.get("cells", [])}
    merges = sheet.get("merged_ranges", [])

    def merge_for(row, column):
        for merged in merges:
            if merged["min_row"] <= row <= merged["max_row"] and merged["min_column"] <= column <= merged["max_column"]:
                return merged
        return None

    anchors = {}
    for selected_cell in selected:
        row, column = selected_cell["row"], selected_cell["column"]
        merged = merge_for(row, column)
        position = (merged["min_row"], merged["min_column"]) if merged else (row, column)
        if position not in cell_by_position:
            raise ValueError(f"Excel display grid is missing selected cell {sheet_name}!{position}")
        role = selected_cell.get("role", "cell")
        if position not in anchors or role in {"header", "row_header"}:
            anchors[position] = role

    source_rows = sorted({row for row, _ in anchors})
    low_column = min(column for _, column in anchors)
    high_column = max(
        (merge_for(row, column) or {"max_column": column})["max_column"]
        for row, column in anchors
    )
    occupied = set()
    output = [f'<table data-sheet="{escape(sheet_name, quote=True)}" data-grid-schema="excel-grid-v1">']
    for row_index, row in enumerate(source_rows):
        output.append(f'<tr data-row="{row}">')
        for column in range(low_column, high_column + 1):
            relative_column = column - low_column
            if (row_index, relative_column) in occupied:
                continue
            role = anchors.get((row, column))
            if role is None:
                output.append("<td></td>")
                continue
            cell = cell_by_position[(row, column)]
            merged = merge_for(row, column)
            width = merged["max_column"] - merged["min_column"] + 1 if merged else 1
            merged_rows = [item for item in source_rows if merged and merged["min_row"] <= item <= merged["max_row"]]
            height = len(merged_rows) if merged else 1
            for covered_row in range(row_index, row_index + height):
                for covered_column in range(relative_column, relative_column + width):
                    occupied.add((covered_row, covered_column))
            tag = "th" if role in {"header", "row_header"} else "td"
            attrs = [f'data-a1="{escape(cell["a1"], quote=True)}"']
            if not compact:
                attrs.extend([
                    f'data-row="{row}"',
                    f'data-column="{column}"',
                    f'data-raw-type="{escape(str(cell.get("raw_type", "unknown")), quote=True)}"',
                ])
            if tag == "th":
                attrs.append('scope="row"' if role == "row_header" else 'scope="col"')
            if height > 1:
                attrs.append(f'rowspan="{height}"')
            if width > 1:
                attrs.append(f'colspan="{width}"')
            display = escape(str(cell.get("display", ""))).replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")
            output.append(f"<{tag} {' '.join(attrs)}>{display}</{tag}>")
        output.append("</tr>")
    output.append("</table>")
    return "".join(output)
def _field(cell):
    return cell["coordinate"] + " " + (cell["header"] + "：" if cell["header"] else "") + cell["value"]

def _prefix(record):
    fields = "; ".join(_field(cell) for cell in record["cells"])
    if not record["text"].endswith(fields):
        raise ValueError("Excel record text does not match its cells")
    return record["text"][:-len(fields)].rstrip() if fields else record["text"]

def _label(header):
    return re.sub(r"[\s_]+", "", header).casefold()

def _source_cell(sheet, cell, role):
    return {
        "sheet": sheet,
        "row": cell["row"],
        "column": cell["column"],
        "a1": cell.get("coordinate") or cell.get("a1") or _a1(cell["row"], cell["column"]),
        "role": role,
    }


def _fragment_source_cells(record, cells):
    source_cells = []
    for cell in cells:
        source_cells.append(_source_cell(record["sheet"], cell, "cell"))
        if cell.get("header") and cell.get("header_row") is not None:
            source_cells.append(
                {
                    "sheet": record["sheet"],
                    "row": cell["header_row"],
                    "column": cell["column"],
                    "a1": _a1(cell["header_row"], cell["column"]),
                    "role": "header",
                }
            )
    for cell in record.get("context_cells", []):
        source_cells.append(_source_cell(record["sheet"], cell, "context"))

    unique = {}
    for cell in source_cells:
        key = (cell["sheet"], cell["row"], cell["column"])
        if key not in unique or cell["role"] == "header":
            unique[key] = cell
    return list(unique.values())


def _row_fragment_groups(record, max_tokens, count_tokens):
    if count_tokens(record["text"]) <= max_tokens:
        return [{"text": record["text"], "source_cells": _fragment_source_cells(record, record["cells"])}]
    cells = record["cells"]
    program = {"프로그램id", "프로그램명", "programid", "programname"}
    definition = {"컬럼명", "컬럼id", "타입", "columnname", "columnid", "type"}
    labels = {_label(cell["header"]) for cell in cells}
    dimensions = program if ({"프로그램id", "프로그램명"} <= labels or {"programid", "programname"} <= labels) else definition if ({"컬럼id", "타입"} <= labels or {"columnid", "type"} <= labels) else set()
    summary_columns = set(record.get("summary_anchor_columns", [])) if record.get("row_role") == "matrix_summary" else set()
    anchors = [cell for cell in cells if _label(cell["header"]) in dimensions or cell["column"] in summary_columns]
    remaining = [cell for cell in cells if cell not in anchors]
    prefix = _prefix(record)
    def render(extra):
        selected = {cell["column"] for cell in anchors + extra}
        selected_cells = [cell for cell in cells if cell["column"] in selected]
        return {
            "text": prefix + " " + "; ".join(_field(cell) for cell in selected_cells),
            "source_cells": _fragment_source_cells(record, selected_cells),
        }
    if not remaining:
        return [{"text": record["text"], "source_cells": _fragment_source_cells(record, cells)}]
    groups, current = [], []
    for cell in remaining:
        proposed = current + [cell]
        if current and count_tokens(render(proposed)["text"]) > max_tokens:
            groups.append(current)
            current = [cell]
        else:
            current = proposed
    if current:
        groups.append(current)
    # Avoid a tiny final fragment by moving complete trailing cells across
    # the preceding boundary, while both fragments remain within the budget.
    minimum = min(64, max_tokens // 4)
    if len(groups) > 1:
        while count_tokens(render(groups[-1])["text"]) < minimum and len(groups[-2]) > 1:
            previous, tail = groups[-2], groups[-1]
            proposed = [previous[-1]] + tail
            if count_tokens(render(proposed)["text"]) > max_tokens:
                break
            groups[-2], groups[-1] = previous[:-1], proposed
    return [render(group) for group in groups]


def _row_fragments(record, max_tokens, count_tokens):
    return [group["text"] for group in _row_fragment_groups(record, max_tokens, count_tokens)]


def _deduplicate_source_cells(source_cells):
    unique = {}
    for cell in source_cells:
        key = (cell["sheet"], cell["row"], cell["column"])
        if key not in unique or cell["role"] == "header":
            unique[key] = cell
    return list(unique.values())


def _logical_table_groups(records):
    tables, current = [], []
    current_sheet = current_context = current_header_row = None
    for record in records:
        context = record.get("context", "")
        header_row = record.get("header_row")
        starts_header = record.get("row_role") == "header" and header_row == record.get("row")
        starts_prelude = record.get("row_role") == "prelude" and header_row != current_header_row
        previous = current[-1] if current else None
        continues_header = starts_header and previous and previous.get("row_role") == "header" and record["row"] == previous["row"] + 1
        same_definition_table = bool(context) and record.get("active_section") in {"index", "script"} and any(
            item.get("active_section") == "definition" for item in current
        )
        boundary = bool(current) and (
            record["sheet"] != current_sheet
            or context != current_context
            or ((starts_header or starts_prelude) and header_row != current_header_row
                and not continues_header and not same_definition_table)
        )
        if boundary:
            tables.append(current)
            current = []
        if not current:
            current_sheet = record["sheet"]
            current_context = context
            current_header_row = header_row
        current.append(record)
    if current:
        tables.append(current)
    return tables


def _table_group(records):
    source_cells = []
    for record in records:
        source_cells.extend(_fragment_source_cells(record, record["cells"]))
    return {"sheet": records[0]["sheet"], "text": "\n".join(record["text"] for record in records),
            "source_cells": _deduplicate_source_cells(source_cells)}


def _whole_table_row_groups(records, max_tokens, count_tokens):
    groups = []
    structural_roles = {"prelude", "header", "section_title"}
    for table in _logical_table_groups(records):
        complete = _table_group(table)
        if count_tokens(complete["text"]) <= max_tokens:
            groups.append(complete)
            continue
        table_context, active_title, active_headers = [], [], []
        current, current_signature, emitted, last_emitted = [], None, set(), []

        def emit(context, body):
            nonlocal last_emitted
            last_emitted = list(context + body)
            group = _table_group(context + body)
            groups.append(group)
            emitted.update(id(record) for record in context + body)

        def emit_structural(records):
            nonlocal last_emitted
            pending = [record for record in records if id(record) not in emitted]
            if not pending:
                return
            appended = _table_group(last_emitted + pending) if last_emitted else None
            if appended and count_tokens(appended["text"]) <= max_tokens:
                groups[-1] = appended
                last_emitted = list(last_emitted + pending)
            else:
                structural = _table_group(table_context + pending)
                if count_tokens(structural["text"]) > max_tokens:
                    raise ValueError("Excel table header context exceeds the configured embedding budget")
                groups.append(structural)
                last_emitted = list(table_context + pending)
            emitted.update(id(record) for record in pending)

        for record in table:
            role = record.get("row_role")
            if role == "prelude":
                table_context.append(record)
                continue
            if role == "section_title":
                if current:
                    emit(current_context, current)
                    current = []
                    current_signature = None
                emit_structural(active_title + active_headers)
                active_title = [record]
                active_headers = []
                continue
            if role == "header":
                active_headers.append(record)
                continue
            context = table_context + active_title + active_headers
            signature = tuple(id(item) for item in context)
            if current and signature != current_signature:
                emit(current_context, current)
                current = []
            candidate = _table_group(context + current + [record])
            if count_tokens(candidate["text"]) > max_tokens:
                if not current:
                    raise ValueError(f"Excel complete row {record['sheet']}!{record['row']} exceeds the configured embedding budget")
                emit(current_context, current)
                current = []
                candidate = _table_group(context + [record])
                if count_tokens(candidate["text"]) > max_tokens:
                    raise ValueError(f"Excel complete row {record['sheet']}!{record['row']} exceeds the configured embedding budget")
            current_context, current_signature = context, signature
            current.append(record)
        if current:
            emit(current_context, current)
        emit_structural([record for record in table if record.get("row_role") in structural_roles])
    return groups

def chunk_excel_record_groups(records, max_tokens=512, count_tokens=None, whole_table_rows=False):
    """Return the unchanged semantic chunk text with explicit source-cell provenance."""
    if max_tokens <= 0:
        raise ValueError("Excel token budget must be positive")
    count_tokens = count_tokens or num_tokens_from_string
    if whole_table_rows:
        return _whole_table_row_groups(records, max_tokens, count_tokens)
    fragments = [
        {"sheet": record["sheet"], **fragment}
        for record in records
        for fragment in _row_fragment_groups(record, max_tokens, count_tokens)
    ]
    groups, current = [], None
    for fragment in fragments:
        proposed_text = current["text"] + "\n" + fragment["text"] if current else fragment["text"]
        if current and count_tokens(proposed_text) > max_tokens:
            current["source_cells"] = _deduplicate_source_cells(current["source_cells"])
            groups.append(current)
            current = {
                "sheet": fragment["sheet"],
                "text": fragment["text"],
                "source_cells": list(fragment["source_cells"]),
            }
        elif current:
            current["text"] = proposed_text
            current["source_cells"].extend(fragment["source_cells"])
        else:
            current = {
                "sheet": fragment["sheet"],
                "text": fragment["text"],
                "source_cells": list(fragment["source_cells"]),
            }
    if current:
        current["source_cells"] = _deduplicate_source_cells(current["source_cells"])
        groups.append(current)
    return groups

def chunk_excel_records(records, max_tokens=512, count_tokens=None):
    """Preserve cell values/context; an indivisible oversize cell stays oversize."""
    return [group["text"] for group in chunk_excel_record_groups(records, max_tokens, count_tokens)]


def attach_excel_display_html(chunks, groups, grid, original_format):
    """Attach HWPX-compatible display HTML after standard chunk tokenization."""
    if len(chunks) != len(groups):
        raise ValueError("Excel display metadata requires one generated group per chunk")
    for chunk, group in zip(chunks, groups, strict=True):
        if chunk.get("content_with_weight") != group["text"]:
            raise ValueError("Excel display metadata cannot alter semantic chunk text")
        metadata = chunk.setdefault("metadata", {}).setdefault("parser_platform", {})
        source_cells = group["source_cells"]
        metadata.update({
            "display_html": render_excel_display_html(grid, source_cells),
            "display_html_only": True,
            "display_html_schema": "excel-grid-v1",
            "source_cells": source_cells,
            "original_format": original_format,
        })
    return chunks
