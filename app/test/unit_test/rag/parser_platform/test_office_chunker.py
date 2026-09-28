from __future__ import annotations

import asyncio
import hashlib
import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from rag.parser_platform.canonical import canonical_sha256
from rag.parser_platform.docling_adapter import DoclingOfficeAdapter
from rag.parser_platform.docling_contract import DoclingOfficeManifest
from rag.parser_platform.office_chunker import OfficeChunker
from rag.parser_platform.schemas import BlockType, DocxProvenance, ParsedBlock, SourceFormat
from rag.parser_platform.standard_bridge import CommonToStandardChunkAdapter
from rag.svr.task_executor_refactor.chunk_service import chunk_parser_platform_document

FIXTURES = Path(__file__).resolve().parents[4] / "test" / "fixtures" / "parser_platform" / "office"


def _document(source_format: SourceFormat):
    if source_format == SourceFormat.XLSX:
        workbook = Workbook()
        first = workbook.active
        first.title = "Sheet-A"
        for row in (("Name", "Value"), ("A", 1), ("B", 2)):
            first.append(row)
        second = workbook.create_sheet("Sheet-B")
        for row_number in (4, 5):
            for column in range(2, 7):
                second.cell(row=row_number, column=column, value=f"v{row_number}_{column}")
        stream = BytesIO()
        workbook.save(stream)
        workbook.close()
        source = stream.getvalue()
    else:
        # The normalizer consumes Docling JSON; source bytes are only needed by
        # the spreadsheet row reconstruction path.
        source = b"synthetic-word-source"
    payload = json.loads((FIXTURES / f"structured.{source_format.value}.docling.json").read_text(encoding="utf-8"))
    manifest = DoclingOfficeManifest(
        task_kind="office_document_parse",
        parse_run_id="test-run",
        source_format=source_format,
        source_hash=hashlib.sha256(source).hexdigest(),
        parser_name="docling",
        parser_version="2.115.0",
        backend="simple-pipeline",
        ocr_enabled=False,
        raw_artifact_hash=canonical_sha256(payload),
        document=payload,
    )
    return DoclingOfficeAdapter().normalize(
        manifest,
        source_document_id="test-document",
        chunk_set_id="test-chunk-set",
        raw_artifact_ref="artifact://test-run/raw.json",
    ), source


@pytest.mark.parametrize("source_format", [SourceFormat.DOC, SourceFormat.DOCX])
def test_word_merges_text_and_preserves_all_locators(source_format: SourceFormat) -> None:
    document, source = _document(SourceFormat.DOCX)
    if source_format == SourceFormat.DOC:
        document = document.model_copy(update={"source_format": SourceFormat.DOC})
    original = CommonToStandardChunkAdapter().adapt(document)
    chunks = OfficeChunker().chunk(document, parser_config={"chunk_token_num": 512, "delimiter": "\n"}, source_bytes=source)
    assert len(chunks) < len(original)
    assert [chunk["chunk_order_int"] for chunk in chunks] == list(range(len(chunks)))
    original_ids = {chunk["metadata"]["parser_platform"]["stable_block_id"] for chunk in original}
    merged_ids = {
        stable_id
        for chunk in chunks
        for stable_id in chunk["metadata"]["parser_platform"]["stable_block_ids"]
    }
    assert merged_ids == original_ids
    for chunk in chunks:
        meta = chunk["metadata"]["parser_platform"]
        assert len(meta["office_locators"]) == len(meta["source_item_ids"])
        assert meta["chunk_token_count"] is not None
        assert meta["contributing_provenance"]


def test_xlsx_row_groups_keep_original_cells_and_sheet_names() -> None:
    document, source = _document(SourceFormat.XLSX)
    chunks = OfficeChunker().chunk(document, parser_config={"excel_chunk_token_num": 128}, source_bytes=source)
    table_chunks = [chunk for chunk in chunks if chunk["doc_type_kwd"] == "table"]
    assert table_chunks
    for chunk in table_chunks:
        meta = chunk["metadata"]["parser_platform"]
        assert meta["source_cells"]
        assert all(cell["sheet"] in {"Sheet-A", "Sheet-B"} for cell in meta["source_cells"])
        assert meta["display_html"].startswith("<table")
        assert meta["office_locators"]
        assert meta["chunk_token_count"] is not None


def test_word_keeps_media_atomic_while_merging_text_across_it() -> None:
    def source(index: int, kind: str, text: str):
        stable_id = f"{index:032x}"
        return {
            "content_with_weight": text,
            "doc_type_kwd": kind,
            "metadata": {"parser_platform": {
                "stable_block_id": stable_id,
                "source_item_id": f"#/texts/{index}",
                "office_locator": {"kind": "docx", "item_locator": f"#/texts/{index}"},
                "contributing_provenance": [{"kind": "docx", "item_locator": f"#/texts/{index}"}],
                "ocr_attachment_ids": ["ocr-one"] if index == 2 else [],
            }},
        }

    chunks = [source(1, "text", "first sentence"), source(2, "image", "image label"), source(3, "text", "second sentence")]
    blocks = [
        SimpleNamespace(
            stable_block_id=f"{index:032x}",
            block_type=BlockType.MEDIA if index == 2 else BlockType.TEXT,
            provenance=(DocxProvenance(item_locator=f"#/texts/{index}"),),
        )
        for index in (1, 2, 3)
    ]
    result = OfficeChunker._word_chunks(
        blocks,
        {chunk["metadata"]["parser_platform"]["stable_block_id"]: chunk for chunk in chunks},
        {"chunk_token_num": 128, "delimiter": "\n"},
    )
    assert [chunk["doc_type_kwd"] for chunk in result] == ["text", "image"]
    merged = result[0]["metadata"]["parser_platform"]
    assert merged["source_item_ids"] == ["#/texts/1", "#/texts/3"]
    assert result[1]["metadata"]["parser_platform"]["ocr_attachment_ids"] == ["ocr-one"]


def test_xlsx_table_ocr_attachment_remains_searchable_once() -> None:
    document, source = _document(SourceFormat.XLSX)
    table = next(block for block in document.blocks if block.block_type == BlockType.TABLE)
    attachment_id = "f" * 32
    attachment = ParsedBlock(
        stable_block_id=attachment_id,
        source_item_id="#/ocr/table",
        block_type=BlockType.OCR_ATTACHMENT,
        reading_order=max(block.reading_order for block in document.blocks) + 1,
        text="OCR TABLE MARKER",
        parent_id=table.stable_block_id,
        provenance=table.provenance,
    )
    document = document.model_copy(update={"blocks": tuple(
        block.model_copy(update={"ocr_attachment_ids": (attachment_id,)}) if block == table else block
        for block in document.blocks
    ) + (attachment,)})
    chunks = OfficeChunker().chunk(document, parser_config={"excel_chunk_token_num": 128}, source_bytes=source)
    matching = [chunk for chunk in chunks if "OCR TABLE MARKER" in chunk["content_with_weight"]]
    assert len(matching) == 1
    meta = matching[0]["metadata"]["parser_platform"]
    assert meta["ocr_attachment_ids"] == [attachment_id]
    assert meta["office_locators"]


@pytest.mark.parametrize("source_format", [SourceFormat.DOC, SourceFormat.DOCX, SourceFormat.XLSX])
def test_task_chunk_dispatch_uses_office_policy(source_format: SourceFormat) -> None:
    document, source = _document(SourceFormat.DOCX if source_format == SourceFormat.DOC else source_format)
    if source_format == SourceFormat.DOC:
        document = document.model_copy(update={"source_format": SourceFormat.DOC})
    chunks = asyncio.run(chunk_parser_platform_document(
        document,
        config=object(),
        parser_config={"chunk_token_num": 512, "delimiter": "\n", "excel_chunk_token_num": 128},
        source_bytes=source,
    ))
    assert chunks
    assert all(chunk["metadata"]["parser_platform"]["chunk_token_count"] is not None for chunk in chunks)
    if source_format != SourceFormat.XLSX:
        assert len(chunks) < len(CommonToStandardChunkAdapter().adapt(document))
    else:
        assert any(chunk["metadata"]["parser_platform"].get("source_cells") for chunk in chunks)
