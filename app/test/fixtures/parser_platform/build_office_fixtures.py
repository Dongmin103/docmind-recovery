from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XlsxImage
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches as PptxInches


ROOT = Path(__file__).parent / "office"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def make_media() -> tuple[Path, Path]:
    ROOT.mkdir(parents=True, exist_ok=True)
    screenshot = ROOT / "screenshot.png"
    image = Image.new("RGB", (640, 320), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((20, 20, 620, 300), outline="navy", width=6)
    draw.text((80, 135), "PROCESS SCREENSHOT OCR TARGET", fill="black")
    image.save(screenshot)

    logo = ROOT / "logo.png"
    icon = Image.new("RGB", (48, 48), "blue")
    ImageDraw.Draw(icon).text((12, 15), "DM", fill="white")
    icon.save(logo)
    return screenshot, logo


def make_docx(screenshot: Path, logo: Path) -> Path:
    output = ROOT / "structured.docx"
    document = Document()
    document.add_heading("품질관리", level=1)
    document.add_heading("시험방법", level=2)
    document.add_paragraph("시험방법 본문은 native text로 보존되어야 합니다.")
    document.add_paragraph("첫 번째 점검 항목", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "시험"
    table.cell(0, 1).text = "결과"
    table.cell(1, 0).text = "함량"
    table.cell(1, 1).text = "적합"
    document.add_picture(str(screenshot), width=Inches(5.0))
    document.add_picture(str(logo), width=Inches(0.35))
    document.save(output)
    return output


def make_xlsx(screenshot: Path) -> Path:
    output = ROOT / "structured.xlsx"
    workbook = Workbook()
    sheet_a = workbook.active
    sheet_a.title = "Sheet-A"
    sheet_a.append(["시험", "결과"])
    sheet_a.append(["함량", 99.5])
    sheet_a.append(["순도", 98.1])
    chart = BarChart()
    chart.add_data(Reference(sheet_a, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(sheet_a, min_col=1, min_row=2, max_row=3))
    sheet_a.add_chart(chart, "D2")

    sheet_b = workbook.create_sheet("Sheet-B")
    sheet_b.merge_cells("B3:F3")
    sheet_b["B3"] = "병합된 시험 결과 영역"
    sheet_b["B4"] = "검체"
    sheet_b["C4"] = "값"
    sheet_b["B5"] = "A"
    sheet_b["C5"] = 10
    screenshot_image = XlsxImage(str(screenshot))
    screenshot_image.width = 320
    screenshot_image.height = 160
    sheet_b.add_image(screenshot_image, "H2")
    workbook.save(output)
    return output


def make_pptx(screenshot: Path) -> Path:
    output = ROOT / "structured.pptx"
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title_slide.shapes.title.text = "공정 검토"
    body = title_slide.shapes.add_textbox(PptxInches(0.8), PptxInches(1.5), PptxInches(4.5), PptxInches(1.0))
    body.text_frame.text = "슬라이드 native body text"
    table = title_slide.shapes.add_table(2, 2, PptxInches(0.8), PptxInches(2.6), PptxInches(4.2), PptxInches(1.5)).table
    table.cell(0, 0).text = "시험"
    table.cell(0, 1).text = "결과"
    table.cell(1, 0).text = "함량"
    table.cell(1, 1).text = "적합"
    chart_data = ChartData()
    chart_data.categories = ["A", "B"]
    chart_data.add_series("값", (10, 20))
    title_slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, PptxInches(5.2), PptxInches(1.5), PptxInches(4.0), PptxInches(3.0), chart_data)

    image_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    image_slide.shapes.add_picture(str(screenshot), PptxInches(0.5), PptxInches(0.5), width=PptxInches(9.0))
    presentation.save(output)
    return output


def main() -> None:
    screenshot, logo = make_media()
    files = [make_docx(screenshot, logo), make_xlsx(screenshot), make_pptx(screenshot)]
    manifest = {
        "schema": "parser-platform-office-fixtures-v1",
        "files": {
            path.name: {"sha256": sha256(path), "bytes": path.stat().st_size}
            for path in [*files, screenshot, logo]
        },
        "required_anchors": {
            "structured.docx": ["품질관리 > 시험방법", "native paragraph", "2x2 table", "screenshot media", "logo media"],
            "structured.xlsx": ["Sheet-A", "Sheet-B!B3:F5", "merged cell", "native chart data", "screenshot media"],
            "structured.pptx": ["slide 1 title/body/table/chart", "slide 2 image-only media"],
        },
    }
    (ROOT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
