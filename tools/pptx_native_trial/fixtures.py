"""Non-sensitive, in-memory PPTX fixtures with independently specified content."""

from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches


def synthetic_deck(slides=1, chart_points=3):
    deck = Presentation()
    for index in range(slides):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(
            Inches(0.3), Inches(0.2), Inches(8), Inches(0.5)
        )
        shape.text = f"시험 슬라이드 {index + 1}"
        body = slide.shapes.add_textbox(Inches(0.3), Inches(0.9), Inches(5), Inches(1))
        body.text = "품질 & 속도 <검증>\n둘째 문단"
        group = slide.shapes.add_group_shape()
        group.shapes.add_textbox(
            Inches(6), Inches(1), Inches(2), Inches(0.4)
        ).text = "그룹 첫째"
        group.shapes.add_textbox(
            Inches(6), Inches(1.5), Inches(2), Inches(0.4)
        ).text = "그룹 둘째"
        table = slide.shapes.add_table(
            3, 3, Inches(0.3), Inches(2), Inches(5), Inches(1.2)
        ).table
        table.cell(0, 0).text = "항목"
        table.cell(0, 1).merge(table.cell(0, 2))
        table.cell(0, 1).text = "병합 제목"
        for row, values in enumerate((("서울", "12", "-3"), ("부산", "0", "8")), 1):
            for column, value in enumerate(values):
                table.cell(row, column).text = value
        data = CategoryChartData()
        data.categories = [f"Q{point + 1}" for point in range(chart_points)]
        data.add_series(
            "매출", [(-3, 0, 12)[point % 3] for point in range(chart_points)]
        )
        data.add_series("비용", [(1, 2, 3)[point % 3] for point in range(chart_points)])
        slide.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(0.3),
            Inches(3.5),
            Inches(5),
            Inches(2),
            data,
        )
    output = BytesIO()
    deck.save(output)
    return output.getvalue()


def rewrite(source, replacements=None, additions=None):
    output = BytesIO()
    with (
        ZipFile(BytesIO(source)) as original,
        ZipFile(output, "w", ZIP_DEFLATED) as result,
    ):
        for entry in original.infolist():
            value = original.read(entry.filename)
            if entry.filename in (replacements or {}):
                value = replacements[entry.filename](value)
            result.writestr(entry.filename, value)
        for name, value in (additions or {}).items():
            result.writestr(name, value)
    return output.getvalue()
