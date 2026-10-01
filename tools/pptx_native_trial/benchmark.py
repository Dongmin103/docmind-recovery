"""Aggregate-only synthetic accuracy/performance trial; no production services."""

from __future__ import annotations

import argparse
import json
import math
import platform
import statistics
import time
from io import BytesIO

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from tools.pptx_native_trial.contracts import load_contracts
from tools.pptx_native_trial.docmind_adapter import to_parsed_document
from tools.pptx_native_trial.evaluation import score
from tools.pptx_native_trial.extractor import VERSION, extract
from tools.pptx_native_trial.fixtures import synthetic_deck


def fixture_accuracy(result, chart_points, *, expected_slides=1):
    """Exact block/cell/tuple F1 against fixture specifications, not parser output."""
    texts, cells, values = [], [], []
    for slide in range(1, expected_slides + 1):
        texts.extend(
            (slide, text)
            for text in (
                f"시험 슬라이드 {slide}",
                "품질 & 속도 <검증>\n둘째 문단",
                "그룹 첫째",
                "그룹 둘째",
            )
        )
        cells.extend(
            (slide, *cell)
            for cell in (
                (0, 0, 1, 1, "항목"),
                (0, 1, 1, 2, "병합 제목"),
                (1, 0, 1, 1, "서울"),
                (1, 1, 1, 1, "12"),
                (1, 2, 1, 1, "-3"),
                (2, 0, 1, 1, "부산"),
                (2, 1, 1, 1, "0"),
                (2, 2, 1, 1, "8"),
            )
        )
        for name, pattern in (("매출", ("-3", "0", "12")), ("비용", ("1", "2", "3"))):
            values.extend(
                (slide, name, f"Q{index + 1}", pattern[index % 3])
                for index in range(chart_points)
            )
    return {
        "text_block_exact": score(
            [(b.slide, b.text) for b in result.blocks if b.kind == "text"], texts
        ),
        "table_cell_and_span_exact": score(
            [(b.slide, *c) for b in result.blocks for c in b.cells], cells
        ),
        "chart_tuple_exact": score(
            [(b.slide, *v) for b in result.blocks for v in b.chart_values], values
        ),
    }


def python_pptx_baseline(source):
    """Library baseline for the same synthetic features, including recursive groups.

    This is deliberately not a full production extractor; no SmartArt or common
    normalization is claimed for this baseline. Arrays are read once per series.
    """
    deck = Presentation(BytesIO(source))
    text_blocks, cells, values = [], [], []
    for number, slide in enumerate(deck.slides, 1):

        def walk(shapes):
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(shape.shapes)
                elif shape.has_text_frame:
                    text_blocks.append((number, shape.text.replace("\v", "\n")))
                elif shape.has_table:
                    for row, items in enumerate(shape.table.rows):
                        for column, cell in enumerate(items.cells):
                            if not cell.is_spanned:
                                cells.append(
                                    (
                                        number,
                                        row,
                                        column,
                                        cell.span_height,
                                        cell.span_width,
                                        cell.text,
                                    )
                                )
                elif shape.has_chart:
                    chart = shape.chart
                    categories = [item.label for item in chart.plots[0].categories]
                    for series in chart.series:
                        series_name = series.name
                        numbers = list(series.values)
                        values.extend(
                            (number, series_name, category, format(value, "g"))
                            for category, value in zip(categories, numbers, strict=True)
                        )

        walk(slide.shapes)
    return text_blocks, cells, values


def timing(function, source, repeats, slides):
    durations = []
    for _ in range(repeats):
        start = time.perf_counter_ns()
        function(source)
        durations.append((time.perf_counter_ns() - start) / 1_000_000)
    ordered = sorted(durations)
    return {
        "repeats": repeats,
        "median_document_ms": round(statistics.median(durations), 3),
        "p95_document_ms": round(ordered[math.ceil(0.95 * repeats) - 1], 3),
        "median_ms_per_slide": round(statistics.median(durations) / slides, 3),
        "p95_ms_per_slide": round(ordered[math.ceil(0.95 * repeats) - 1] / slides, 3),
        "max_ms_per_slide": round(max(durations) / slides, 3),
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repeats", type=int, default=10)
    parser.add_argument("--slides", type=int, nargs="+", default=[1, 25, 100])
    args = parser.parse_args()
    if not 1 <= args.repeats <= 100 or any(
        not 1 <= size <= 500 for size in args.slides
    ):
        parser.error("repeats must be 1..100; slides must be 1..500")
    load_contracts()
    from rag.parser_platform.chunk_adapter import CommonToStandardChunkAdapter

    def normalized(source):
        result = extract(source)
        parsed = to_parsed_document(result)
        CommonToStandardChunkAdapter().adapt(parsed)
        return result

    scenarios = [(size, 3) for size in args.slides] + [(1, 1000)]
    report = {
        "parser_version": VERSION,
        "python": platform.python_version(),
        "platform": platform.platform(),
        "scope": "synthetic PPTX only; no OCR, embedding, indexing, decryption or real-document claims",
        "normalization": "real DocMind schema + common chunk-field adapter, not slide token chunking",
        "performance_boundary": "in-memory source bytes through output; ZIP/XML/hash included; fixtures built outside timer",
        "percentiles": "document duration divided by slides, not individual-slide latency; fresh extraction each repetition",
        "scenarios": [],
    }
    for slides, chart_points in scenarios:
        source = synthetic_deck(slides, chart_points)
        first_start = time.perf_counter_ns()
        result = normalized(source)
        first_ms = (time.perf_counter_ns() - first_start) / 1_000_000
        baseline = python_pptx_baseline(source)
        expected = (
            [(b.slide, b.text) for b in result.blocks if b.kind == "text"],
            [(b.slide, *c) for b in result.blocks for c in b.cells],
            [(b.slide, *v) for b in result.blocks for v in b.chart_values],
        )
        accuracy = fixture_accuracy(result, chart_points, expected_slides=slides)
        scenario = {
            "slides": slides,
            "chart_points_per_series": chart_points,
            "source_bytes": len(source),
            "xml_bytes_read": result.xml_bytes,
            "blocks": len(result.blocks),
            "warnings": result.warnings,
            "native_coverage_complete": result.coverage_complete,
            "first_normalized_ms_per_slide": round(first_ms / slides, 3),
            "native_core": timing(extract, source, args.repeats, slides),
            "native_with_common_normalization": timing(
                normalized, source, args.repeats, slides
            ),
            "python_pptx_core": timing(
                python_pptx_baseline, source, args.repeats, slides
            ),
            "synthetic_accuracy": accuracy,
            "python_pptx_agrees_with_native": all(
                score(a, b)["f1"] == 1 for a, b in zip(baseline, expected, strict=True)
            ),
        }
        scenario["trial_gate_pass"] = (
            result.slide_count == slides
            and not result.warnings
            and all(m["f1"] == 1 for m in accuracy.values())
            and scenario["native_with_common_normalization"]["p95_ms_per_slide"] <= 50
        )
        report["scenarios"].append(scenario)
    print(json.dumps(report, ensure_ascii=True, indent=2))
    return 0 if all(item["trial_gate_pass"] for item in report["scenarios"]) else 1


if __name__ == "__main__":
    raise SystemExit(main())
